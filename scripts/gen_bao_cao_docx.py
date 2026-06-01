"""Tạo file Word báo cáo đồ án AIFAS từ mô tả cố định (đồng bộ với codebase)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm
from docx.oxml.ns import qn


def set_doc_defaults(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Cm(2.5)
    sec.bottom_margin = Cm(2.5)
    sec.left_margin = Cm(3)
    sec.right_margin = Cm(2)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style.font.size = Pt(13)


def add_title_block(doc: Document) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BÁO CÁO MÔN HỌC\nCHUYÊN ĐỀ CƠ SỞ DỮ LIỆU")
    r.bold = True
    r.font.size = Pt(14)
    doc.add_paragraph()
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run(
        "XÂY DỰNG HỆ THỐNG AIFAS\n"
        "(AI Fear Analytics System)\n"
        "Tích hợp kho dữ liệu, OLAP/Data Mining SSAS và ứng dụng Web dự đoán mức độ lo ngại về AI"
    )
    r2.bold = True
    r2.font.size = Pt(16)
    doc.add_paragraph()
    for line in ("Mã đề tài / nhóm: ……………………………", "Giảng viên hướng dẫn: ……………………………", "Sinh viên thực hiện: ……………………………"):
        p3 = doc.add_paragraph(line)
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    p4 = doc.add_paragraph("Tháng 5/2026")
    p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()


def main() -> None:
    doc = Document()
    set_doc_defaults(doc)
    add_title_block(doc)

    doc.add_heading("Lời cảm ơn", level=1)
    doc.add_paragraph(
        "Nhóm em xin gửi lời cảm ơn tới … (điền tên giảng viên, bạn bè, gia đình) "
        "đã hỗ trợ trong quá trình thực hiện đồ án."
    )

    doc.add_heading("Mục lục", level=1)
    doc.add_paragraph(
        "Trong Microsoft Word: References → Table of Contents → Automatic Table 1 "
        "(sau khi chỉnh sửa nội dung, bấm chuột phải mục lục → Update Field)."
    )

    doc.add_heading("Danh mục chữ viết tắt", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    hdr[0].text = "Ký hiệu"
    hdr[1].text = "Ý nghĩa"
    rows = [
        ("SSAS", "SQL Server Analysis Services"),
        ("OLAP", "Online Analytical Processing"),
        ("ETL", "Extract – Transform – Load"),
        ("DMX", "Data Mining Extensions (ngôn ngữ truy vấn mô hình khai phá)"),
        ("MDX", "Multidimensional Expressions"),
        ("API", "Application Programming Interface"),
        ("AIFAS", "AI Fear Analytics System — hệ thống phân tích khảo sát lo ngại về AI"),
    ]
    for a, b in rows:
        row = table.add_row().cells
        row[0].text = a
        row[1].text = b

    doc.add_heading("Chương 1. Tổng quan đề tài", level=1)

    doc.add_heading("1.1. Đặt vấn đề", level=2)
    doc.add_paragraph(
        "Sự phổ biến của trí tuệ nhân tạo (AI) đi kèm nhiều lo ngại của người dùng về "
        "việc mất việc làm, quyền riêng tư, an toàn và tác động xã hội. Khảo sát thu thập "
        "các biến nhân khẩu học, bối cảnh sử dụng AI và các chỉ số sợ hãi (thang điểm) "
        "cần được lưu trữ có cấu trúc, tổng hợp đa chiều (OLAP) và khai phá (Data Mining) "
        "để hỗ trợ dự đoán mức độ lo ngại tổng quát (Fear Level: High / Medium / Low)."
    )

    doc.add_heading("1.2. Mục tiêu hệ thống", level=2)
    doc.add_paragraph(
        "Xây dựng chuỗi xử lý dữ liệu hoàn chỉnh: (1) CSDL nguồn lưu khảo sát thô; "
        "(2) mô hình sao (star schema) phục vụ OLAP; (3) CSDL khai phá AIFAS_DM; "
        "(4) dự án SSAS với mô hình khai phá; (5) gói ETL (SSIS) nếu triển khai tự động; "
        "(6) ứng dụng Web ASP.NET Core phục vụ giao diện và API gọi DMX tới SSAS."
    )

    doc.add_heading("1.3. Phạm vi và công nghệ", level=2)
    doc.add_paragraph(
        "Microsoft SQL Server (Database Engine), SQL Server Analysis Services (đa chiều / "
        "data mining tùy cấu hình dự án), Integration Services (ETL), .NET 8 (Minimal API), "
        "thư viện Microsoft.AnalysisServices.AdomdClient, giao diện tĩnh HTML/CSS/JS."
    )

    doc.add_heading("Chương 2. Phân tích và thiết kế dữ liệu", level=1)

    doc.add_heading("2.1. CSDL nguồn AIFAS_SOURCE", level=2)
    doc.add_paragraph(
        "File AIFAS_SOURCE.sql định nghĩa database AIFAS_SOURCE và bảng AI_FEAR_SURVEY "
        "chứa dữ liệu khảo sát gốc: Age, Gender, Education, Occupation, AI_Familiarity, "
        "Job_Loss_Fear, Privacy_Fear, Safety_Fear, Social_Impact_Fear, Overall_Fear_Score, "
        "AI_Usage_Frequency, Country, Income_Level, FearLevel. Script cấp quyền db_datareader "
        "cho tài khoản dịch vụ Analysis Services (NT SERVICE\\MSOLAP$…) để SSAS đọc dữ liệu."
    )

    doc.add_heading("2.2. Mô hình sao AIFAS_STAR", level=2)
    doc.add_paragraph(
        "File AIFAS_STAR.sql mô tả kho dữ liệu dạng sao gồm các chiều DimPerson "
        "(PersonID, Age, Gender, Education, Occupation), DimContext (AI_Familiarity, "
        "AI_Usage_Frequency, Country, Income_Level), DimTime (SurveyYear, SurveyQuarter, "
        "SurveyMonth) và bảng sự kiện FactAIFear liên kết khóa ngoại tới ba chiều, lưu các "
        "chỉ số sợ hãi và FearLevel. Thiết kế này phù hợp truy vấn OLAP theo nhiều trục phân tích."
    )

    doc.add_heading("2.3. CSDL phục vụ khai phá AIFAS_DM", level=2)
    doc.add_paragraph(
        "File SQLQuery3.sql tạo AIFAS_DM, bảng DATASET có cấu trúc tương đương nguồn, "
        "sao chép dữ liệu từ AIFAS_SOURCE (INSERT … SELECT), bổ sung cột ID_INT nếu cần "
        "khóa số nguyên, và cấp quyền đọc cho SSAS trên AIFAS_DM, AIFAS_SOURCE, AIFAS_STAR."
    )

    doc.add_heading("Chương 3. OLAP và Data Mining (SSAS)", level=1)
    doc.add_heading("3.1. Vai trò của SSAS", level=2)
    doc.add_paragraph(
        "Thư mục AIFAS_SSAS chứa solution/dự án Analysis Services (file .dwproj, .dsv, .dmm…). "
        "Dữ liệu huấn luyện mô hình khai phá thường trỏ tới AIFAS_DM (hoặc view nguồn tương đương). "
        "Sau khi deploy, catalog SSAS (ví dụ tên AIFAS_SSAS trong appsettings) chứa mining structure "
        "và các mining model."
    )

    doc.add_heading("3.2. Các mô hình khai phá trong codebase ứng dụng", level=2)
    doc.add_paragraph(
        "Ứng dụng Web (Program.cs) gọi DMX tới ba model: DECISION_TREE và LOGISTIC_REGRESSION "
        "dự đoán cột mục tiêu [Fear Level] với xác suất High/Medium/Low; CLUSTERING dùng "
        "Cluster() và ClusterProbability() để gán nhóm và độ tin cậy. Các cột đầu vào trong "
        "chuỗi NATURAL PREDICTION JOIN khớp mining structure (tên có khoảng trắng như "
        "[AI Familiarity], [Overall Fear Score]). Lưu ý: mô hình DECISION_TREE trong project "
        "không gồm cột Overall Fear Score nên API chỉ truyền trường này cho LOGISTIC_REGRESSION "
        "và CLUSTERING (includeOverall=true/false trong mã nguồn)."
    )

    doc.add_heading("3.3. Truy vấn MDX/DMX tham khảo", level=2)
    doc.add_paragraph(
        "File AIFAS_DM.mdx chứa ví dụ kiểm tra metadata ($system.DMSCHEMA_MINING_MODELS, "
        "DMSCHEMA_MINING_COLUMNS), xem nội dung cây DECISION_TREE.CONTENT và các câu DMX "
        "PredictProbability / Cluster cho từng model — phục vụ kiểm thử trong SSMS."
    )

    doc.add_heading("Chương 4. ETL (tùy chọn triển khai)", level=1)
    doc.add_paragraph(
        "Thư mục ETL_AIFAS chứa gói SSIS (Package.dtsx) có thể dùng để tự động hóa luồng "
        "nạp/chuyển đổi dữ liệu giữa các database. Chi tiết mapping cần mô tả theo cấu hình "
        "thực tế khi mở Visual Studio / SSDT."
    )

    doc.add_heading("Chương 5. Ứng dụng Web và API", level=1)

    doc.add_heading("5.1. Kiến trúc", level=2)
    doc.add_paragraph(
        "Dự án AIFAS_WebApi là ứng dụng ASP.NET Core (TargetFramework net8.0), Minimal API không "
        "controller riêng. StaticFiles và DefaultFiles trỏ tới thư mục cha của project (repo root) "
        "để phục vụ các trang HTML tĩnh (trangchu.html làm trang mặc định)."
    )

    doc.add_heading("5.2. Kết nối SSAS", level=2)
    doc.add_paragraph(
        "Chuỗi kết nối lấy từ biến môi trường SSAS_CONNECTION_STRING hoặc appsettings.json "
        "mục Ssas:ConnectionString (Data Source, Initial Catalog=AIFAS_SSAS, Integrated Security…). "
        "File SSAS_KET_NOI.txt hướng dẫn xử lý lỗi kết nối: bật dịch vụ Analysis Services, "
        "đúng instance (ví dụ …\\\\SSASMD), kiểm tra bằng SSMS, có thể ghi đè chuỗi bằng PowerShell "
        "trước khi chạy dotnet run."
    )

    doc.add_heading("5.3. API", level=2)
    api_tbl = doc.add_table(rows=1, cols=3)
    api_tbl.style = "Table Grid"
    h = api_tbl.rows[0].cells
    h[0].text = "Phương thức / đường dẫn"
    h[1].text = "Mô tả"
    h[2].text = "Đầu ra"
    for path, desc, out in [
        ("GET /api/health", "Kiểm tra API sống", "JSON { ok: true }"),
        (
            "POST /api/predict",
            "Nhận JSON PredictionRequest (Age, Gender, Education, Occupation, AiFamiliarity, "
            "AiUsageFrequency, Country, IncomeLevel, JobLossFear, PrivacyFear, SafetyFear, "
            "SocialImpactFear, OverallFearScore)",
            "JSON PredictResponse: kết quả DecisionTree, LogisticRegression (nhãn + xác suất), "
            "Clustering (ClusterLabel, Probability)",
        ),
    ]:
        c = api_tbl.add_row().cells
        c[0].text = path
        c[1].text = desc
        c[2].text = out

    doc.add_heading("5.4. Giao diện người dùng", level=2)
    doc.add_paragraph(
        "Các trang HTML (trangchu.html, phantich.html, ungdung.html, baiviet*.html, lienhe.html) "
        "dùng bố cục thống nhất (font Syne / Be Vietnam Pro, tông màu tối, accent xanh lá). "
        "Trang ungdung.html chứa biểu mẫu nhập thông tin khảo sát và gọi fetch('/api/predict', …) "
        "để hiển thị kết quả dự đoán từ SSAS."
    )

    doc.add_heading("Chương 6. Hướng dẫn cài đặt và chạy thử", level=1)
    steps = [
        "Cài SQL Server (Engine), SSAS, tùy chọn SSIS; tạo và nạp dữ liệu theo các script .sql.",
        "Chạy / deploy dự án SSAS, xác nhận tên catalog và mining model khớp mã (DECISION_TREE, LOGISTIC_REGRESSION, CLUSTERING).",
        "Cập nhật appsettings.json (Ssas:ConnectionString) cho đúng Data Source và Initial Catalog.",
        "Từ thư mục AIFAS_WebApi: dotnet run — mở trình duyệt tại URL hiển thị (thường http://localhost:5xxx).",
        "Vào trang Ứng dụng, điền form và bấm dự đoán; nếu lỗi 502, đọc thông báo JSON và đối chiếu SSAS_KET_NOI.txt.",
    ]
    for i, s in enumerate(steps, 1):
        doc.add_paragraph(f"{i}. {s}", style="List Number")

    doc.add_heading("Chương 7. Kết luận và hướng phát triển", level=1)
    doc.add_paragraph(
        "Hệ thống AIFAS minh họa đầy đủ pipeline: nguồn → kho (star) → vùng khai phá → "
        "mô hình SSAS → API hiện đại → giao diện Web. Hướng mở rộng: xác thực người dùng, "
        "logging, đóng gói Docker, thay HTML tĩnh bằng SPA, đánh giá độ chính xác mô hình "
        "trên tập kiểm định, và bổ sung dashboard OLAP (Excel/Power BI) trên AIFAS_STAR."
    )

    doc.add_heading("Tài liệu tham khảo", level=1)
    doc.add_paragraph(
        "[1] Microsoft Learn — SQL Server Documentation.\n"
        "[2] Microsoft Learn — Analysis Services / DMX syntax.\n"
        "[3] Tài liệu giảng dạy môn Chuyên đề CSDL — … (điền theo giáo trình trường)."
    )

    doc.add_heading("Phụ lục: Cấu trúc thư mục mã nguồn (tham khảo)", level=1)
    doc.add_paragraph(
        "ChuyenDe_CSDL/\n"
        "  AIFAS_SOURCE.sql, AIFAS_STAR.sql, SQLQuery3.sql — script SQL\n"
        "  AIFAS_DM.mdx — truy vấn MDX/DMX mẫu\n"
        "  AIFAS_SSAS/ — solution Analysis Services\n"
        "  ETL_AIFAS/ — gói SSIS\n"
        "  AIFAS_WebApi/ — Program.cs, appsettings.json, SSAS_KET_NOI.txt\n"
        "  *.html — giao diện tĩnh ở root repo"
    )

    out = Path(__file__).resolve().parent.parent / "BaoCao_AIFAS_ChuyenDeCSDL.docx"
    doc.save(out)
    print("Created:", str(out))  # ASCII: tránh lỗi encoding console Windows


if __name__ == "__main__":
    main()
